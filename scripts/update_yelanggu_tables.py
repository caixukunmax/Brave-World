import openpyxl
import os

DATA_DIR = os.path.join(os.path.dirname(__file__), '..', 'tables', 'datas', 'common')
ITEM_DIR = os.path.join(os.path.dirname(__file__), '..', 'tables', 'datas', 'item')

MAP_CONFIG_FILE = os.path.join(DATA_DIR, '#MapConfig-地图配置表.xlsx')
MONSTER_FILE = os.path.join(DATA_DIR, '#Monster-怪物表.xlsx')
DROP_GROUP_FILE = os.path.join(DATA_DIR, '#DropGroup-掉落组表.xlsx')
MAP_MONSTER_FILE = os.path.join(DATA_DIR, '#MapMonster-地图怪物表.xlsx')

MAP_ID = 5
MAP_NAME = '野狼谷'
MAP_WIDTH = 80
MAP_HEIGHT = 80
SPAWN_X = 40
SPAWN_Y = 40

MONSTER_WOLF_KING_ID = 12
MONSTER_SILVER_WOLF_ID = 13

# 野狼：12 个刷新点（常见）
WOLF_SPAWNS = [
    (18, 18), (24, 14), (32, 22), (14, 32), (62, 18), (68, 24),
    (58, 32), (16, 62), (24, 56), (32, 66), (62, 64), (56, 56)
]

# 野狼王：3 个刷新点（北部岩石高地）
KING_SPAWNS = [
    (36, 10), (44, 12), (48, 8)
]

# 白银野狼：1 个刷新点（深处密林核心，最稀有）
SILVER_SPAWNS = [
    (52, 52)
]


def find_last_data_row(ws, id_col=2):
    """找到数据最后一行（id 列非空）"""
    last = ws.max_row
    while last > 1 and ws.cell(last, id_col).value is None:
        last -= 1
    return last


def update_map_config():
    wb = openpyxl.load_workbook(MAP_CONFIG_FILE)
    ws = wb.active
    last = find_last_data_row(ws)

    # 检查是否已存在
    for r in range(5, last + 1):
        if ws.cell(r, 2).value == MAP_ID:
            print(f'[MapConfig] 野狼谷 (id={MAP_ID}) 已存在，跳过')
            wb.close()
            return

    row = last + 1
    ws.cell(row, 2, MAP_ID)
    ws.cell(row, 3, MAP_NAME)
    ws.cell(row, 4, MAP_NAME)
    ws.cell(row, 5, MAP_WIDTH)
    ws.cell(row, 6, MAP_HEIGHT)
    ws.cell(row, 7, SPAWN_X)
    ws.cell(row, 8, SPAWN_Y)

    wb.save(MAP_CONFIG_FILE)
    wb.close()
    print(f'[MapConfig] 新增 id={MAP_ID} {MAP_NAME}')


def update_monsters():
    wb = openpyxl.load_workbook(MONSTER_FILE)
    ws = wb.active
    last = find_last_data_row(ws)

    existing_ids = set()
    for r in range(5, last + 1):
        existing_ids.add(ws.cell(r, 2).value)

    monsters = [
        {
            'id': MONSTER_WOLF_KING_ID,
            'name': '野狼王',
            'level': 8,
            'exp': 100,
            'attrs': 'HP=180|ATK=18|DEF=6|AGILITY=115',
            'skills': '撕咬|狼嚎|猛毒撕咬|寒冰吐息',
            'drop_group_id': 12,
        },
        {
            'id': MONSTER_SILVER_WOLF_ID,
            'name': '白银野狼',
            'level': 6,
            'exp': 80,
            'attrs': 'HP=110|ATK=12|DEF=4|AGILITY=140',
            'skills': '撕咬|猛毒撕咬',
            'drop_group_id': 13,
        },
    ]

    for m in monsters:
        if m['id'] in existing_ids:
            print(f'[Monster] {m["name"]} (id={m["id"]}) 已存在，跳过')
            continue
        row = last + 1
        ws.cell(row, 2, m['id'])
        ws.cell(row, 3, m['name'])
        ws.cell(row, 4, m['level'])
        ws.cell(row, 5, m['exp'])
        # F=drop_items 留空
        ws.cell(row, 7, m['attrs'])
        ws.cell(row, 8, m['skills'])
        ws.cell(row, 9, m['drop_group_id'])
        last = row
        print(f'[Monster] 新增 {m["name"]} (id={m["id"]})')

    wb.save(MONSTER_FILE)
    wb.close()


def update_drop_groups():
    wb = openpyxl.load_workbook(DROP_GROUP_FILE)
    ws = wb.active
    last = find_last_data_row(ws)

    existing_ids = set()
    for r in range(5, last + 1):
        existing_ids.add(ws.cell(r, 2).value)

    groups = [
        {
            'id': 12,
            'name': '野狼王掉落',
            'entries': '1002,1,3,80,true|2001,1,1,30,false|1006,1,1,5,false',
        },
        {
            'id': 13,
            'name': '白银野狼掉落',
            'entries': '1002,1,2,60,true|2001,1,1,20,false|1006,1,1,15,false',
        },
    ]

    for g in groups:
        if g['id'] in existing_ids:
            print(f'[DropGroup] id={g["id"]} 已存在，跳过')
            continue
        row = last + 1
        ws.cell(row, 2, g['id'])
        ws.cell(row, 3, g['name'])
        ws.cell(row, 4, g['entries'])
        last = row
        print(f'[DropGroup] 新增 id={g["id"]} {g["name"]}')

    wb.save(DROP_GROUP_FILE)
    wb.close()


def update_map_monsters():
    wb = openpyxl.load_workbook(MAP_MONSTER_FILE)
    ws = wb.active
    last = find_last_data_row(ws)
    last_id = int(ws.cell(last, 2).value or 0)

    # 检查是否已存在 map_id=5 的记录
    existing_count = 0
    for r in range(5, last + 1):
        if ws.cell(r, 3).value == MAP_ID:
            existing_count += 1
    if existing_count > 0:
        print(f'[MapMonster] map_id={MAP_ID} 已有 {existing_count} 条记录，跳过')
        wb.close()
        return

    spawns = []
    # 野狼：常见，12 个点，复活 45 秒
    for x, y in WOLF_SPAWNS:
        spawns.append({
            'monster_id': 3,
            'x': x,
            'y': y,
            'respawn_time': 45,
            'respawn_range': 3,
        })
    # 野狼王：精英，3 个点，复活 120 秒
    for x, y in KING_SPAWNS:
        spawns.append({
            'monster_id': MONSTER_WOLF_KING_ID,
            'x': x,
            'y': y,
            'respawn_time': 120,
            'respawn_range': 2,
        })
    # 白银野狼：稀有，1 个点，复活 300 秒
    for x, y in SILVER_SPAWNS:
        spawns.append({
            'monster_id': MONSTER_SILVER_WOLF_ID,
            'x': x,
            'y': y,
            'respawn_time': 300,
            'respawn_range': 1,
        })

    for i, s in enumerate(spawns):
        row = last + 1 + i
        last_id += 1
        ws.cell(row, 2, last_id)
        ws.cell(row, 3, MAP_ID)
        ws.cell(row, 4, s['monster_id'])
        ws.cell(row, 5, s['x'])
        ws.cell(row, 6, s['y'])
        ws.cell(row, 7, s['respawn_time'])
        ws.cell(row, 8, True)
        ws.cell(row, 9, 2)            # ai_id=2 patrol_chase
        ws.cell(row, 10, 'SpawnPoint')
        ws.cell(row, 11, s['respawn_range'])

    wb.save(MAP_MONSTER_FILE)
    wb.close()
    print(f'[MapMonster] 新增 {len(spawns)} 个刷新点 for map_id={MAP_ID}')


if __name__ == '__main__':
    update_map_config()
    update_monsters()
    update_drop_groups()
    update_map_monsters()
