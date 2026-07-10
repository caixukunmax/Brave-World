import openpyxl
import os

DATA_DIR = 'C:/code/tslua2/tables/datas/common'
MAP_CONFIG_FILE = os.path.join(DATA_DIR, '#MapConfig-地图配置表.xlsx')
MAP_MONSTER_FILE = os.path.join(DATA_DIR, '#MapMonster-地图怪物表.xlsx')

# 1. 更新 MapConfig，添加落叶乡
wb_config = openpyxl.load_workbook(MAP_CONFIG_FILE)
ws_config = wb_config.active

# 找到最后一行数据（用 id 列 B 判断）
last_row = ws_config.max_row
while last_row > 1 and ws_config.cell(last_row, 2).value is None:
    last_row -= 1

# 添加新行
new_id = 4
ws_config.cell(last_row + 1, 2, new_id)
ws_config.cell(last_row + 1, 3, '落叶乡')
ws_config.cell(last_row + 1, 4, '落叶乡')
ws_config.cell(last_row + 1, 5, 100)
ws_config.cell(last_row + 1, 6, 100)
ws_config.cell(last_row + 1, 7, 50)
ws_config.cell(last_row + 1, 8, 50)

wb_config.save(MAP_CONFIG_FILE)
print(f'[AddMonsters] MapConfig added 落叶乡 with id={new_id}')

# 2. 更新 MapMonster，添加 3 只哥布林
wb_monster = openpyxl.load_workbook(MAP_MONSTER_FILE)
ws_monster = wb_monster.active

last_row = ws_monster.max_row
while last_row > 1 and ws_monster.cell(last_row, 2).value is None:
    last_row -= 1

last_id = int(ws_monster.cell(last_row, 2).value or 0)
spawns = [
    (55, 50),
    (50, 55),
    (45, 50),
]

for i, (x, y) in enumerate(spawns):
    row = last_row + 1 + i
    new_monster_id = last_id + 1 + i
    ws_monster.cell(row, 2, new_monster_id)
    ws_monster.cell(row, 3, new_id)        # map_id
    ws_monster.cell(row, 4, 2)             # monster_id 哥布林
    ws_monster.cell(row, 5, x)
    ws_monster.cell(row, 6, y)
    ws_monster.cell(row, 7, 30)            # respawn_time
    ws_monster.cell(row, 8, True)          # is_active
    ws_monster.cell(row, 9, 2)             # ai_id
    ws_monster.cell(row, 10, 'SpawnPoint') # respawn_type
    ws_monster.cell(row, 11, 0)            # respawn_range

wb_monster.save(MAP_MONSTER_FILE)
print(f'[AddMonsters] MapMonster added {len(spawns)} goblins for map_id={new_id}')
