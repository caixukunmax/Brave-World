import openpyxl
import os

DATA_DIR = os.path.join(os.path.dirname(__file__), '..', 'tables', 'datas', 'common')

MAP_CONFIG_FILE = os.path.join(DATA_DIR, '#MapConfig-地图配置表.xlsx')
MAP_MONSTER_FILE = os.path.join(DATA_DIR, '#MapMonster-地图怪物表.xlsx')


def remove_map_config():
    wb = openpyxl.load_workbook(MAP_CONFIG_FILE)
    ws = wb.active
    removed = False
    # 数据从第 5 行开始
    for r in range(ws.max_row, 4, -1):
        if ws.cell(r, 2).value == 1:
            ws.delete_rows(r)
            removed = True
            print(f'[MapConfig] 删除 id=1 新手村 (行 {r})')
    wb.save(MAP_CONFIG_FILE)
    wb.close()
    if not removed:
        print('[MapConfig] 未找到 id=1 新手村')


def remove_map_monsters():
    wb = openpyxl.load_workbook(MAP_MONSTER_FILE)
    ws = wb.active
    removed = 0
    for r in range(ws.max_row, 4, -1):
        if ws.cell(r, 3).value == 1:
            ws.delete_rows(r)
            removed += 1
    wb.save(MAP_MONSTER_FILE)
    wb.close()
    print(f'[MapMonster] 删除 map_id=1 的 {removed} 条记录')


if __name__ == '__main__':
    remove_map_config()
    remove_map_monsters()
