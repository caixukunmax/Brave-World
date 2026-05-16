import openpyxl
from openpyxl.styles import Font, PatternFill

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'TerrainConfig'

# Row 1: ##var + field names
headers = ['##var', 'id', 'name', 'description', 'walkable', 'move_speed_ratio', 'can_swim',
    'patk_modifier', 'matk_modifier', 'pdef_modifier', 'mdef_modifier',
    'hp_regen_per_sec', 'mp_regen_per_sec', 'fire_damage_bonus', 'ice_damage_bonus',
    'poison_damage_bonus', 'color_r', 'color_g', 'color_b', 'color_a', 'particle_effect']
ws.append(headers)

# Row 2: ##type + types
types = ['##type', 'int', 'string', 'string', 'bool', 'float', 'bool',
    'float', 'float', 'float', 'float',
    'int', 'int', 'int', 'int', 'int',
    'int', 'int', 'int', 'float', 'string']
ws.append(types)

# Row 3: ##group (optional)
groups = ['##group'] + [''] * 20
ws.append(groups)

# Row 4: ## + comments
comments = ['##', '地形ID', '名称', '描述', '可行走', '移速倍率', '可游泳',
    '物攻修正', '魔攻修正', '物防修正', '魔防修正',
    'HP恢复/秒', 'MP恢复/秒', '火伤加成', '冰伤加成', '毒伤加成',
    '颜色R', '颜色G', '颜色B', '透明度', '粒子路径']
ws.append(comments)

# Data rows - first column (A) must be empty/None, data starts from B
# This matches existing Luban xlsx format
terrains = [
    [None, 0, '普通', '平坦地形', 'true', 1.0, 'false', 1.0, 1.0, 1.0, 1.0, 0, 0, 0, 0, 0, 200, 200, 200, 0.0, ''],
    [None, 1, '水域', '深水区域', 'false', 0.0, 'true', 0.8, 1.2, 0.5, 1.0, 0, 2, 0, 5, 0, 50, 100, 200, 0.3, ''],
    [None, 2, '草地', '茂密草地', 'true', 1.0, 'false', 1.0, 1.0, 1.0, 1.0, 1, 0, 0, 0, 2, 50, 180, 50, 0.2, ''],
    [None, 3, '沙地', '松软沙地', 'true', 0.7, 'false', 1.0, 0.9, 0.8, 0.9, 0, 0, 0, 0, 0, 200, 180, 80, 0.2, ''],
    [None, 4, '岩石', '坚硬岩石', 'true', 1.0, 'false', 1.1, 0.9, 1.2, 1.0, 0, 0, 0, 0, 0, 120, 120, 120, 0.3, ''],
    [None, 5, '雪地', '积雪覆盖', 'true', 0.6, 'false', 1.0, 1.1, 1.0, 0.8, 0, 0, 0, 5, 0, 240, 240, 255, 0.25, ''],
    [None, 6, '沼泽', '泥泞沼泽', 'true', 0.4, 'false', 0.9, 1.0, 0.7, 0.7, -2, 0, 0, 0, 3, 80, 100, 60, 0.35, ''],
    [None, 7, '岩浆', '灼热岩浆', 'false', 0.0, 'false', 1.2, 1.0, 1.0, 0.5, -5, 0, 10, 0, 0, 200, 60, 20, 0.4, ''],
    [None, 8, '神圣地', '神圣领域', 'true', 1.0, 'false', 1.0, 1.2, 1.0, 1.2, 5, 5, 0, 0, 0, 255, 255, 200, 0.2, ''],
]

for t in terrains:
    ws.append(t)

# Style header rows
for row_idx in [1, 2, 3, 4]:
    for cell in ws[row_idx]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='E0E0E0', end_color='E0E0E0', fill_type='solid')

# Adjust column widths
ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 6
ws.column_dimensions['C'].width = 10
ws.column_dimensions['D'].width = 12
for col_letter in ['E','F','G','H','I','J','K','L','M','N','O','P']:
    ws.column_dimensions[col_letter].width = 14
for col_letter in ['Q','R','S','T']:
    ws.column_dimensions[col_letter].width = 10
ws.column_dimensions['U'].width = 20

# Save
output_path = 'C:/code/tslua2/tables/datas/common/TbTerrainConfig.xlsx'
wb.save(output_path)
print(f'Saved: {output_path}')
