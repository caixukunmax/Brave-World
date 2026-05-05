import openpyxl
wb = openpyxl.load_workbook(r'C:\code\tslua2\tables\datas\__beans__.xlsx')
ws = wb.active
for row in range(1, ws.max_row + 1):
    v2 = ws.cell(row, 2).value
    if v2 and 'Buff' in str(v2):
        print(f'Row {row}: full_name={v2}')
        # Print all fields for this bean
        r = row
        while r <= ws.max_row:
            name = ws.cell(r, 10).value
            typ = ws.cell(r, 12).value
            if name and typ:
                print(f'  field: {name} type={typ}')
            elif ws.cell(r, 2).value and r > row:
                break
            r += 1