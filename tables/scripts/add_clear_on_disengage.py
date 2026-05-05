import openpyxl

# 1. Add field to __beans__.xlsx (BuffConfigRow definition)
wb = openpyxl.load_workbook(r'C:\code\tslua2\tables\datas\__beans__.xlsx')
ws = wb.active

# Find the row after shield_base in BuffConfigRow
# BuffConfigRow starts at row 44, shield_base is at some row after that
in_buff = False
shield_row = None
for row in range(44, ws.max_row + 1):
    v2 = ws.cell(row, 2).value
    if v2 and v2 != 'common.BuffConfigRow':
        # We've hit the next bean definition
        break
    name = ws.cell(row, 10).value
    if name == 'shield_base':
        shield_row = row
    if v2 == 'common.BuffConfigRow':
        in_buff = True

print(f'shield_base at row {shield_row}')

# Insert clear_on_disengage after shield_base
# We need to add a new row. Shift everything down first.
# Actually, just find the next empty row after shield_base within BuffConfigRow
next_row = shield_row + 1
# Check if next row is empty or belongs to next bean
next_v2 = ws.cell(next_row, 2).value
next_name = ws.cell(next_row, 10).value
if next_v2 or next_name:
    # Need to insert a row
    ws.insert_rows(next_row)

ws.cell(next_row, 10, 'clear_on_disengage')
ws.cell(next_row, 12, 'bool')
ws.cell(next_row, 14, '脱战时是否清除')
wb.save(r'C:\code\tslua2\tables\datas\__beans__.xlsx')
print(f'Added clear_on_disengage field at row {next_row}')

# 2. Update the data xlsx - add column with default true
wb2 = openpyxl.load_workbook(r'C:\code\tslua2\tables\datas\common\#Buff-Buff表.xlsx')
ws2 = wb2.active
# clear_on_disengage is already at col 13 from earlier, but let's verify
col13_header = ws2.cell(1, 13).value
if col13_header == 'clear_on_disengage':
    print('Data xlsx already has clear_on_disengage column')
else:
    ws2.cell(1, 13, 'clear_on_disengage')
    for row in range(2, ws2.max_row + 1):
        ws2.cell(row, 13, True)
    print('Added clear_on_disengage column to data xlsx')
wb2.save(r'C:\code\tslua2\tables\datas\common\#Buff-Buff表.xlsx')
