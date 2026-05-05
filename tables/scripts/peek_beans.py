import openpyxl
wb = openpyxl.load_workbook(r'C:\code\tslua2\tables\datas\__beans__.xlsx')
ws = wb.active
for row in range(1, min(ws.max_row + 1, 30)):
    vals = []
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row, col).value
        if v: vals.append(f'[{col}]={v}')
    if vals:
        print(f'Row {row}: {" ".join(vals)}')
